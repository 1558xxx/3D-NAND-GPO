"""Build complete per-chip/per-experiment state cohorts from migrated raw matrices."""
from pathlib import Path
import json,hashlib,time
import numpy as np
import openpyxl
from scipy.io import loadmat
HERE=Path(__file__).resolve().parent

def atomic_json(path,value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix(path.suffix+'.tmp');temp.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    for attempt in range(12):
        try:temp.replace(path);return
        except PermissionError:time.sleep(.15*(attempt+1))
    # A transient Windows indexer lock on replacement must not abort completed numerical work.
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')

def sha(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda:f.read(1048576),b''):h.update(b)
    return h.hexdigest()

def main():
    out=HERE/'data';out.mkdir(exist_ok=True);records=[]
    previous={}
    if (out/'progress.json').exists():
        saved=json.loads((out/'progress.json').read_text(encoding='utf-8'))
        previous={(r['chip'],r['experiment'],r['state']):r for r in saved.get('records',[])}
    if (out/'manifest.json').exists():print('Data already prepared',flush=True);return
    for chip,vendor in [('TH58TFT1T23BA8H','Kioxia CT'),('MT29F1T08EEHAF','Micron FG')]:
      for group,levels,unit in [('Bake',[1,10,100,1000,10000],'equivalent_minutes'),('Read',[1,1000,2000,3000,4000,5000],'read_repetitions')]:
        if all((chip,group,k) in previous and (out/previous[(chip,group,k)]['file']).exists() for k in range(1,8)):
            records.extend(previous[(chip,group,k)] for k in range(1,8));continue
        folder=HERE/'raw'/f'{chip}_{group}';mat=folder/'wlcount.mat'
        arr=loadmat(mat)['count'];nlevels=len(levels);scan=131 if chip.startswith('TH') else 164
        assert arr.ndim==3 and arr.shape[1:]==(nlevels*8,11)
        assert arr.shape[0]%scan==0
        nwl=arr.shape[0]//scan
        assert np.isfinite(arr).all() and (arr>=0).all()
        workbook=folder/(chip+'.xlsx')
        w=openpyxl.load_workbook(workbook,read_only=True,data_only=True)
        first=next(w.worksheets[0].iter_rows(min_row=1,max_row=1,max_col=1,values_only=True))[0]
        assert first=='wl:0'
        pseudo=scan+8
        initial=list(w.worksheets[0].iter_rows(min_row=1,max_row=pseudo,max_col=8,values_only=True))
        steps=np.array([i for i,r in enumerate(initial) if isinstance(r[0],(float,int))],dtype=np.float32)
        assert len(steps)==scan
        assert w.sheetnames==['PE-0']+[f'PE-{i*1000}' for i in range(1,11)]
        w.close()
        # Rows: PEC, within-family stress level, WL. Source/target separation by index remains explicit.
        wl=np.tile(np.arange(nwl,dtype=np.int16),nlevels*11)
        label=np.tile(np.repeat(np.arange(1,nlevels+1,dtype=np.int8),nwl),11)
        pec=np.repeat(np.arange(0,11000,1000,dtype=np.int16),nlevels*nwl)
        physical=np.asarray(levels,dtype=np.float32)[label-1]
        task_id=np.arange(len(wl),dtype=np.int32)
        for state in range(1,8):
            if (chip,group,state) in previous and (out/previous[(chip,group,state)]['file']).exists():
                records.append(previous[(chip,group,state)]);continue
            values=arr[:,np.arange(nlevels)*8+state,:].reshape(nwl,scan,nlevels,11).transpose(3,2,0,1).reshape(-1,scan)
            lengths=(values>0).sum(1)
            valid=lengths>=12;keep=np.flatnonzero(valid);maxlen=int(lengths[valid].max())
            ys=np.zeros((len(keep),maxlen),np.float32);xs=np.zeros_like(ys);mask=np.zeros_like(ys,dtype=bool)
            for row,k in enumerate(keep):
                nonzero=np.flatnonzero(values[k]>0);length=len(nonzero)
                xs[row,:length]=steps[nonzero];ys[row,:length]=values[k,nonzero];mask[row,:length]=True
            path=out/f'{chip}_{group}_P{state}.npz'
            np.savez_compressed(path,steps=xs,frequency=ys,mask=mask,WL=wl[keep],label=label[keep],PEC=pec[keep],
                physical=physical[keep],task_id=task_id[keep],source=label[keep]<=3,full_scan_steps=steps)
            row=dict(chip=chip,vendor=vendor,experiment=group,unit=unit,state=state,wordlines=int(nwl),scan_count=scan,
                source_curves=int(np.sum(label[keep]<=3)),target_curves=int(np.sum(label[keep]>3)),total_curves=len(keep),
                excluded_short_curves=int((~valid).sum()),all_candidate_curves=len(valid),min_curve_points=int(lengths[valid].min()),
                max_curve_points=maxlen,file=path.name,raw_mat_sha256=sha(mat),data_sha256=sha(path),
                source_labels=[1,2,3],target_labels=list(range(4,nlevels+1)),stress_values=levels,
                preprocessing='all valid positive-count curves; no cohort subsampling; min12 is the existing protocol rule',
                coordinate='original within-WL worksheet positions including separator gaps; not absolute voltage')
            records.append(row);atomic_json(out/'progress.json',dict(status='preparing',finished_states=len(records),expected_states=28,records=records))
            print(f'{chip} {group} P{state}: source {row["source_curves"]}, target {row["target_curves"]}, excluded {row["excluded_short_curves"]}',flush=True)
        del arr
    atomic_json(out/'manifest.json',dict(status='complete',records=records,source_curves=sum(r['source_curves'] for r in records),
                target_curves=sum(r['target_curves'] for r in records),family_separation=True,zero_counts_kept_in_raw=True,
                baseline_shared_between_families='first-level acquisition is reused; independent target domains are evaluated separately'))
    print('Complete raw-data cohort preparation',flush=True)

if __name__=='__main__':main()
